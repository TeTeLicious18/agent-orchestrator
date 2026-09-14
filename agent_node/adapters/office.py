"""Spreadsheet tools for the autonomous adapter.

Excel files are edited through openpyxl rather than by driving the Excel UI: the
result is deterministic, works whether or not Excel is installed, and does not depend
on window focus. Once written, ``open_file`` shows the workbook in Excel if the operator
wants to see it.

Formulas are written as ordinary cell values - ``=SUM(A1:A10)`` becomes a live formula
when Excel opens the file. openpyxl does not evaluate them, so ``excel_read`` returns
the formula text, not a computed number.
"""

from __future__ import annotations

import asyncio
import os
from typing import Any

from .base import AdapterError

MAX_ROWS = 5_000
MAX_COLUMNS = 256


def _require_openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError:
        raise AdapterError(
            "spreadsheet tools require openpyxl: pip install -r requirements-office.txt"
        ) from None
    return openpyxl


class ExcelController:
    """All paths arrive already resolved inside the workspace by the adapter."""

    async def write(self, path: str, sheet: str | None, rows: Any, append: bool = False) -> dict[str, Any]:
        if not isinstance(rows, list) or not rows:
            raise AdapterError("'rows' must be a non-empty list of lists")
        if len(rows) > MAX_ROWS:
            raise AdapterError(f"at most {MAX_ROWS} rows per call")

        normalised: list[list[Any]] = []
        for row in rows:
            if not isinstance(row, list):
                raise AdapterError("every entry in 'rows' must itself be a list of cell values")
            if len(row) > MAX_COLUMNS:
                raise AdapterError(f"at most {MAX_COLUMNS} columns per row")
            normalised.append(row)

        def _work() -> dict[str, Any]:
            openpyxl = _require_openpyxl()
            exists = os.path.isfile(path)
            book = openpyxl.load_workbook(path) if exists else openpyxl.Workbook()

            if sheet:
                page = book[sheet] if sheet in book.sheetnames else book.create_sheet(sheet)
            else:
                page = book.active

            if not append:
                page.delete_rows(1, page.max_row)
            for row in normalised:
                page.append(row)

            os.makedirs(os.path.dirname(path), exist_ok=True)
            book.save(path)
            return {"ok": True, "sheet": page.title, "rows_written": len(normalised), "total_rows": page.max_row}

        return await asyncio.to_thread(_work)

    async def read(self, path: str, sheet: str | None, max_rows: int = 200) -> dict[str, Any]:
        def _work() -> dict[str, Any]:
            openpyxl = _require_openpyxl()
            if not os.path.isfile(path):
                raise AdapterError("workbook not found")
            book = openpyxl.load_workbook(path)
            page = book[sheet] if sheet and sheet in book.sheetnames else book.active
            # Cap at the real extent: iter_rows would otherwise pad with empty rows and
            # flood the model's context.
            limit = min(max_rows, MAX_ROWS, page.max_row)
            rows = [
                ["" if cell is None else cell for cell in row]
                for row in page.iter_rows(max_row=limit, values_only=True)
            ]
            while rows and all(value == "" for value in rows[-1]):
                rows.pop()
            return {
                "ok": True,
                "sheet": page.title,
                "sheets": book.sheetnames,
                "rows": rows,
                "truncated": page.max_row > limit,
            }

        return await asyncio.to_thread(_work)

    async def set_cell(self, path: str, sheet: str | None, cell: Any, value: Any) -> dict[str, Any]:
        if not isinstance(cell, str) or not cell.strip():
            raise AdapterError("'cell' is required, e.g. 'B4'")

        def _work() -> dict[str, Any]:
            openpyxl = _require_openpyxl()
            if not os.path.isfile(path):
                raise AdapterError("workbook not found; create it with excel_write first")
            book = openpyxl.load_workbook(path)
            page = book[sheet] if sheet and sheet in book.sheetnames else book.active
            page[cell.strip().upper()] = value
            book.save(path)
            return {"ok": True, "sheet": page.title, "cell": cell.strip().upper(), "value": value}

        return await asyncio.to_thread(_work)
